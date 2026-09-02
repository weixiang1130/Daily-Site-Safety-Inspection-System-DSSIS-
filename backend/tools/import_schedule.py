# -*- coding: utf-8 -*-
"""把工地的列控表（MS Project 匯出）匯入地端資料庫，供「今日重點工項」使用。

資料來源
--------
各工地的「七合一列控表」Excel 內含一張 `project` 工作表，是 MS Project 的
任務匯出：識別碼／名稱／工期／開始時間／完成時間／大綱階層。
本工具只讀這張表，其他工作表（甘特圖畫面、備註）不碰。

只匯入葉節點：大綱階層中沒有子項的才是實際作業；上層是「基礎工程」
這類彙總，掛上牆沒有意義。

整份替換：列控表每月改版，工項名稱會改、識別碼會重排，逐筆比對新舊
版本沒有可靠的鍵。因此每次匯入先刪掉同一工地、同一來源的舊資料再寫入。

用法
----
    python backend/tools/import_schedule.py <列控表.xlsx> --site <工地代碼>

    例：python backend/tools/import_schedule.py 列控表.xlsx --site BD04
        python backend/tools/import_schedule.py 列控表.xlsx --site BD05 --dry-run

列控表更新時重跑一次即可。檔案路徑與工地代碼由參數提供，
本工具不寫死任何工地名稱或內部路徑（本 repo 公開）。
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from datetime import date, datetime

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__))), "onprem"))

try:
    import openpyxl
except ImportError:
    sys.exit("需要 openpyxl，請先執行：pip install openpyxl")

# MS Project 中文版匯出的日期格式："2024年9月12日 上午 08:00"
RE_DATE = re.compile(r"(\d{4})年(\d{1,2})月(\d{1,2})日")


def parse_date(v) -> date | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    m = RE_DATE.search(str(v))
    return date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def as_int(v) -> int:
    try:
        return int(str(v).strip())
    except (ValueError, TypeError):
        return 0


def read_tasks(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if "project" not in wb.sheetnames:
        sys.exit(f"找不到 project 工作表（實際有：{wb.sheetnames}）。\n"
                 "本工具讀的是 MS Project 匯出的任務表，請確認列控表格式。")

    rows = [r for r in wb["project"].iter_rows(min_row=2, values_only=True)
            if r[3]]      # 第 4 欄是名稱，空白列跳過

    tasks = []
    for i, r in enumerate(rows):
        rid, _active, _mode, name, _dur, st, en, _pred, lvl, _note = r[:10]
        s, e, level = parse_date(st), parse_date(en), as_int(lvl)
        if not s or not e:
            continue
        # 葉節點 = 下一列的大綱階層沒有更深（沒有子項）
        nxt = as_int(rows[i + 1][8]) if i + 1 < len(rows) else 0
        tasks.append({
            "task_no": str(rid) if rid is not None else "",
            "name": str(name).strip(),
            "start": s, "end": e,
            "level": level, "leaf": nxt <= level,
        })
    return tasks


def main() -> None:
    ap = argparse.ArgumentParser(description="匯入列控表工項")
    ap.add_argument("xlsx", help="列控表檔案路徑")
    ap.add_argument("--site", required=True, help="工地代碼（如 BD04）")
    ap.add_argument("--dry-run", action="store_true", help="只列出將匯入的內容")
    args = ap.parse_args()

    tasks = read_tasks(args.xlsx)
    leafs = [t for t in tasks if t["leaf"]]
    today = date.today()
    active = [t for t in leafs if t["start"] <= today <= t["end"]]

    print(f"解析 {os.path.basename(args.xlsx)}：全部 {len(tasks)} 項、"
          f"葉工項 {len(leafs)} 項、今日進行中 {len(active)} 項")
    for t in active:
        print(f"  今日進行中：[{t['task_no']}] {t['name']} "
              f"{t['start']}~{t['end']}")

    if args.dry_run:
        print("（dry-run，未寫入）")
        return

    from app.db import PlannedTask, SessionLocal, Site, init_db
    init_db()
    db = SessionLocal()
    try:
        site = db.query(Site).filter(Site.code == args.site).first()
        if not site:
            sys.exit(f"工地代碼 {args.site} 不存在。先跑一次表單同步"
                     "（collectors.sync_forms）讓工地主檔進來，或確認代碼。")

        old = (db.query(PlannedTask)
               .filter(PlannedTask.site_code == args.site,
                       PlannedTask.source == "schedule").delete())
        for t in leafs:
            db.add(PlannedTask(
                site_id=site.id, site_code=args.site, source="schedule",
                task_no=t["task_no"], name=t["name"],
                start_date=t["start"], end_date=t["end"],
                outline_level=t["level"], is_leaf=True,
            ))
        db.commit()
        print(f"已寫入 {site.name}（{args.site}）：替換舊資料 {old} 筆 → "
              f"新資料 {len(leafs)} 筆")
    finally:
        db.close()


if __name__ == "__main__":
    main()
